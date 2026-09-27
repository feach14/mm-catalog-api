#!/usr/bin/env python3
"""Import legacy material catalog Excel exports through Catalog API HTTP endpoints."""

import argparse
import csv
import json
import mimetypes
import ssl
import sys
import time
import urllib.error
import urllib.request
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


COUNT_TYPES = {1: "M2", 2: "PM", 3: "SHT", 4: "PERCENT", 5: "LIST", 6: "PACK", 7: "FIX", 8: "KG"}
DELETED_SUFFIX = " (удалён)"


class ApiError(RuntimeError):
    def __init__(self, method, path, status, body):
        super().__init__(f"{method} {path}: HTTP {status}: {body[:500]}")
        self.status = status
        self.body = body


class Client:
    def __init__(self, catalog_url, accounts_url, insecure):
        self.catalog_url = catalog_url.rstrip("/")
        self.accounts_url = accounts_url.rstrip("/")
        self.context = ssl._create_unverified_context() if insecure else ssl.create_default_context()
        self.cookie = None

    def request(self, method, path, data=None, *, accounts=False, multipart=None):
        base = self.accounts_url if accounts else self.catalog_url
        headers = {}
        body = None
        if self.cookie:
            headers["Cookie"] = self.cookie
        if data is not None:
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        if multipart is not None:
            filename, content, content_type = multipart
            boundary = "----CatalogImport" + uuid.uuid4().hex
            body = (
                f"--{boundary}\r\n"
                f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
                f"Content-Type: {content_type}\r\n\r\n"
            ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
            headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
        request = urllib.request.Request(base + path, data=body, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, context=self.context, timeout=60) as response:
                payload = response.read()
                return json.loads(payload.decode("utf-8")) if payload else None
        except urllib.error.HTTPError as error:
            text = error.read().decode("utf-8", errors="replace")
            raise ApiError(method, path, error.code, text) from error

    def sign_in(self, phone, password):
        request = urllib.request.Request(
            self.accounts_url + "/api/sign/in",
            data=json.dumps({"phone": phone, "password": password, "isManager": True}).encode(),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, context=self.context, timeout=60) as response:
            raw_cookie = response.headers.get("Set-Cookie")
            if not raw_cookie or "AccountsApiCookie=" not in raw_cookie:
                raise RuntimeError("Accounts API did not return AccountsApiCookie")
            self.cookie = raw_cookie.split(";", 1)[0]


def read_sheet(path):
    worksheet = load_workbook(path, read_only=False, data_only=True)["Result 1"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, row)) for row in rows[1:]]


def image_bytes(value):
    if isinstance(value, bytes):
        return value
    if not isinstance(value, str) or not value.startswith("0x"):
        raise ValueError("Image data is not a SQL hexadecimal value")
    return bytes.fromhex(value[2:])


def normalized_image(row):
    content = image_bytes(row["Data"])
    content_type = row["Type"]
    if content.startswith(b"BM"):
        source = Image.open(BytesIO(content))
        output = BytesIO()
        source.save(output, format="PNG")
        return output.getvalue(), "image/png", ".png"
    extension = mimetypes.guess_extension(content_type) or ".bin"
    return content, content_type, extension


def material_payload(row, category_id, image_guid):
    return {
        "categoryId": category_id,
        "name": row["Name"],
        "article": row["Article"],
        "image": image_guid,
        "count": 0 if row["Price"] <= 0 else row["Count"],
        "size": row["Size"],
        "depth": row["Depth"],
        "kvM": row["KvM"],
        "perimetrM": row["PerimetrM"],
        "applicableToRaskroys": row["ApplicableToRaskroys"],
        "applicableToPvhFacades": row["ApplicableToPvhFacades"],
        "applicableToEmalFacades": row["ApplicableToEmalFacades"],
        "commentOnMaterialIsRequired": row["CommentOnMaterialIsRequired"],
        "allowSecondItemInOrder": row["AllowSecondItemInOrder"],
        "externalLink": row["ExternalLink"],
        "price": row["Price"],
        "countTypeEnum": COUNT_TYPES[row["CountTypeEnum"]],
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--catalog-url", default="https://127.0.0.1:5005")
    parser.add_argument("--accounts-url", default="https://127.0.0.1:5003")
    parser.add_argument("--phone", required=True)
    parser.add_argument("--password", required=True)
    parser.add_argument("--insecure", action="store_true")
    parser.add_argument("--state", type=Path, default=Path("/tmp/mm_catalog_import_state.json"))
    parser.add_argument("--id-map", type=Path, default=Path(__file__).parent / "docs" / "MATERIAL_IMPORT_ID_MAP.csv")
    args = parser.parse_args()

    categories = read_sheet(args.source / "MaterialCategories.xlsx")
    materials = read_sheet(args.source / "Materials.xlsx")
    images = read_sheet(args.source / "MaterialImages.xlsx")
    images_by_material = {row["MaterialId"]: row for row in images}

    client = Client(args.catalog_url, args.accounts_url, args.insecure)
    client.sign_in(args.phone, args.password)
    state = {"categories": {}, "materials": {}, "pendingImages": {}, "failures": []}
    if args.state.exists():
        state.update(json.loads(args.state.read_text()))

    current_categories = client.request("GET", "/api/materials/categories")["items"]
    category_by_name = {row["name"]: row for row in current_categories}
    for row in sorted(categories, key=lambda value: value["OrderByCol"]):
        existing = category_by_name.get(row["Name"])
        if existing:
            new_id = existing["id"]
        else:
            new_id = client.request("POST", "/api/materials/categories", {
                "name": row["Name"], "externalLink": row["ExternalLink"]
            })["id"]
        state["categories"][str(row["Id"])] = new_id
        args.state.write_text(json.dumps(state, ensure_ascii=False, indent=2))
    print(f"Categories ready: {len(state['categories'])}", flush=True)

    corrected_zero_price = 0
    materials_by_old_id = {str(row["Id"]): row for row in materials}
    for old_id, new_id in list(state["materials"].items()):
        row = materials_by_old_id[old_id]
        if row["Price"] > 0:
            continue
        card = client.request("GET", f"/api/materials/admin/{new_id}")
        payload = material_payload(row, state["categories"][str(row["CategoryId"])], card.get("image"))
        client.request("PUT", f"/api/materials/{new_id}", payload)
        corrected_zero_price += 1
    if corrected_zero_price:
        print(f"Zero-price materials corrected to count=0: {corrected_zero_price}", flush=True)

    current_materials = client.request(
        "GET", "/api/materials/admin?raskroy=true&pvhFacades=true&emalFacades=true"
    )["materials"]
    material_by_name = {row["name"]: row for row in current_materials}

    ordered_materials = sorted(materials, key=lambda value: (
        next(c["OrderByCol"] for c in categories if c["Id"] == value["CategoryId"]),
        value["OrderByCol"], value["Id"]
    ))
    for index, row in enumerate(ordered_materials, 1):
        old_id = str(row["Id"])
        if old_id in state["materials"]:
            continue
        existing = material_by_name.get(row["Name"])
        if existing:
            state["materials"][old_id] = existing["id"]
            state["pendingImages"].pop(old_id, None)
            args.state.write_text(json.dumps(state, ensure_ascii=False, indent=2))
            continue
        image_guid = state["pendingImages"].get(old_id)
        image = images_by_material.get(row["Id"])
        try:
            if image and not image_guid:
                content, content_type, extension = normalized_image(image)
                uploaded = client.request(
                    "POST", "/api/materials/images",
                    multipart=(f"material-{row['Id']}{extension}", content, content_type),
                )
                image_guid = uploaded["fileGuid"]
                state["pendingImages"][old_id] = image_guid
                args.state.write_text(json.dumps(state, ensure_ascii=False, indent=2))
            payload = material_payload(row, state["categories"][str(row["CategoryId"])], image_guid)
            if row["Deleted"] and payload["name"].endswith(DELETED_SUFFIX):
                payload["name"] = payload["name"][:-len(DELETED_SUFFIX)]
            new_id = client.request("POST", "/api/materials", payload)["id"]
            if row["Deleted"]:
                client.request("DELETE", f"/api/materials/{new_id}")
            state["materials"][old_id] = new_id
            state["pendingImages"].pop(old_id, None)
            state["failures"] = [failure for failure in state["failures"] if failure["oldId"] != row["Id"]]
        except Exception as error:
            state["failures"].append({"oldId": row["Id"], "name": row["Name"], "error": str(error)})
            print(f"FAILED old id {row['Id']}: {error}", file=sys.stderr, flush=True)
        args.state.write_text(json.dumps(state, ensure_ascii=False, indent=2))
        if index % 25 == 0 or index == len(ordered_materials):
            print(f"Materials processed: {index}/{len(ordered_materials)}; imported={len(state['materials'])}; failures={len(state['failures'])}", flush=True)

    actual_categories = client.request("GET", "/api/materials/categories")["items"]
    actual_materials = client.request("GET", "/api/materials/admin?raskroy=true&pvhFacades=true&emalFacades=true")["materials"]
    expected_active = [row for row in materials if not row["Deleted"] and str(row["Id"]) in state["materials"]]
    report = {
        "source": {"categories": len(categories), "materials": len(materials), "images": len(images),
                   "activeMaterials": sum(not row["Deleted"] for row in materials),
                   "deletedMaterials": sum(bool(row["Deleted"]) for row in materials)},
        "imported": {"categories": len(state["categories"]), "materials": len(state["materials"]),
                     "activeMaterialsVisible": len(actual_materials)},
        "verification": {
            "categoryCountMatches": len(actual_categories) == len(categories),
            "activeMaterialCountMatchesImported": len(actual_materials) == len(expected_active),
        },
        "failures": state["failures"],
        "stateFile": str(args.state),
        "idMapFile": str(args.id_map),
    }
    failures_by_id = {failure["oldId"]: failure["error"] for failure in state["failures"]}
    args.id_map.parent.mkdir(parents=True, exist_ok=True)
    with args.id_map.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(["EntityType", "OldId", "NewId", "Name", "Status", "Error"])
        for row in sorted(categories, key=lambda value: value["Id"]):
            new_id = state["categories"].get(str(row["Id"]))
            writer.writerow(["MaterialCategory", row["Id"], new_id, row["Name"], "Imported" if new_id else "Failed", ""])
        for row in sorted(materials, key=lambda value: value["Id"]):
            new_id = state["materials"].get(str(row["Id"]))
            writer.writerow(["Material", row["Id"], new_id, row["Name"], "Imported" if new_id else "Failed", failures_by_id.get(row["Id"], "")])
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not state["failures"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
